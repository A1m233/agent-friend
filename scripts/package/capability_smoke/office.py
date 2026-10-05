"""Exercise Office wheels inside managed Python, never the bridge's interpreter."""

from __future__ import annotations

import json
import subprocess
from pathlib import Path

# Kept as source because these libraries belong to managed Python, not PyInstaller.
OFFICE_SCRIPT = r"""
import hashlib
import json
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pandas as pd
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


book = Workbook()
sheet = book.active
sheet.title = "Sales"
sheet.append(["item", "amount"])
sheet.append(["苹果", 12])
sheet.append(["香蕉", 8])
sheet["B4"] = "=SUM(B2:B3)"
book.create_sheet("Notes")["A1"] = "说明"
book.save("source.xlsx")
book.close()

# A stale cache deliberately differs from the current inputs; no calculation engine runs.
namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
with zipfile.ZipFile("source.xlsx") as original:
    entries = [(item, original.read(item)) for item in original.infolist()]
with zipfile.ZipFile("source.xlsx", "w", zipfile.ZIP_DEFLATED) as output:
    for item, data in entries:
        if item.filename == "xl/worksheets/sheet1.xml":
            xml = ET.fromstring(data)
            xml.find("s:sheetData/s:row/s:c[@r='B4']/s:v", namespace).text = "99"
            data = ET.tostring(xml, encoding="utf-8")
        output.writestr(item, data)
source_hash = digest("source.xlsx")

formulas = load_workbook("source.xlsx", read_only=True, data_only=False)
cached = load_workbook("source.xlsx", read_only=True, data_only=True)
assert formulas.sheetnames == ["Sales", "Notes"]
assert list(formulas["Sales"].iter_rows(min_row=2, max_row=3, max_col=2, values_only=True)) == [("苹果", 12), ("香蕉", 8)]
assert formulas["Sales"]["B4"].value == "=SUM(B2:B3)"
assert cached["Sales"]["B4"].value == 99
formulas.close()
cached.close()
frame = pd.read_excel("source.xlsx", sheet_name="Sales", nrows=2, engine="openpyxl")
assert int(frame["amount"].sum()) == 20
assert digest("source.xlsx") == source_hash

edited = load_workbook("source.xlsx")
edited["Sales"]["B2"] = 15
edited.save("edited.xlsx")
edited.close()
reopened = load_workbook("edited.xlsx")
assert reopened["Sales"]["B2"].value == 15
assert reopened["Sales"]["B4"].value == "=SUM(B2:B3)"
reopened.close()
reopened = load_workbook("edited.xlsx", data_only=True)
assert reopened["Sales"]["B4"].value is None
reopened.close()
assert digest("source.xlsx") == source_hash

for filename, sep in [("input.csv", ","), ("input.tsv", "\t")]:
    # Include quoting, Unicode and embedded newlines as well as numbers.
    expected = pd.DataFrame({"item": ["苹果,红", "香蕉\n黄"], "amount": [12, 8]})
    expected.to_csv(filename, sep=sep, index=False, encoding="utf-8")
    before = digest(filename)
    actual = pd.read_csv(filename, sep=sep, encoding="utf-8")
    pd.testing.assert_frame_equal(actual, expected)
    assert digest(filename) == before
    actual["amount"] = actual["amount"] + 1
    actual.to_excel(filename + ".xlsx", index=False, engine="openpyxl")
    pd.testing.assert_frame_equal(pd.read_excel(filename + ".xlsx", engine="openpyxl"), actual)

document = Document()
document.add_heading("季度报告", level=1)
document.add_paragraph("读取正文与表格。")
table = document.add_table(rows=2, cols=2)
for row, values in zip(table.rows, [("项目", "数量"), ("苹果", "12")]):
    for cell, text in zip(row.cells, values):
        cell.text = text
document.save("source.docx")
doc_hash = digest("source.docx")
read = Document("source.docx")
assert [p.text for p in read.paragraphs] == ["季度报告", "读取正文与表格。"]
assert [[c.text for c in row.cells] for row in read.tables[0].rows] == [["项目", "数量"], ["苹果", "12"]]
assert digest("source.docx") == doc_hash
read.paragraphs[1].text = "已更新正文。"
read.tables[0].cell(1, 1).text = "15"
read.save("edited.docx")
reopened = Document("edited.docx")
assert reopened.paragraphs[1].text == "已更新正文。"
assert reopened.tables[0].cell(1, 1).text == "15"
assert digest("source.docx") == doc_hash

rejected = []
for name, reader, error in [
    ("broken.xlsx", load_workbook, zipfile.BadZipFile),
    ("broken.docx", Document, PackageNotFoundError),
    ("legacy.xls", load_workbook, InvalidFileException),
]:
    Path(name).write_bytes(b"not an Office package")
    try:
        reader(name)
    except error:
        rejected.append(name)
    else:
        raise AssertionError("invalid Office input was accepted: " + name)

print(json.dumps({
    "xlsx": True, "csv_tsv": True, "docx": True,
    "formula_cache_distinguished": True, "originals_unchanged": True,
    "invalid_inputs_rejected": rejected,
}, sort_keys=True))
"""


def run_office_smoke(python: Path, environment: dict[str, str], workdir: Path) -> dict[str, object]:
    """Fail closed on a missing wheel, native DLL, template or incorrect file result."""
    completed = subprocess.run(
        [str(python), "-I", "-B", "-c", OFFICE_SCRIPT],
        env=environment,
        cwd=workdir,
        stdin=subprocess.DEVNULL,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=60,
        check=False,
    )
    if completed.returncode:
        raise RuntimeError(f"managed Office smoke failed: {completed.stderr[-4000:]}")
    expected: dict[str, object] = {
        "xlsx": True,
        "csv_tsv": True,
        "docx": True,
        "formula_cache_distinguished": True,
        "originals_unchanged": True,
        "invalid_inputs_rejected": ["broken.xlsx", "broken.docx", "legacy.xls"],
    }
    if json.loads(completed.stdout) != expected:
        raise RuntimeError("managed Office smoke returned unexpected results")
    return expected
